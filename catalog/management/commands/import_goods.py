"""Load the original legal goods wording from the approved MISTERCHIP XLSX."""
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from catalog.models import Category, TrademarkGood
from catalog.seed import CATEGORY_SLUGS


class Command(BaseCommand):
    help = "Import the 373 verbatim Nice class 9 entries from the approved Excel catalogue"

    def add_arguments(self, parser):
        parser.add_argument("xlsx_file")
        parser.add_argument("--check-only", action="store_true", help="Validate without database writes")

    def handle(self, *args, **options):
        book = load_workbook(options["xlsx_file"], read_only=True, data_only=True)
        if "Каталог товаров" not in book.sheetnames:
            raise CommandError("В файле отсутствует лист «Каталог товаров».")
        sheet = book["Каталог товаров"]
        rows = list(sheet.iter_rows(values_only=True))
        header_at = next((i for i, row in enumerate(rows)
                          if len(row) >= 3 and str(row[0]).strip() == "№ в перечне свидетельства"
                          and str(row[1]).strip() == "Категория интернет-магазина"
                          and str(row[2]).strip() == "Товар (дословно по свидетельству)"), None)
        if header_at is None:
            raise CommandError("Не найдены точные названия трёх исходных колонок.")
        goods = []
        for row in rows[header_at + 1:]:
            if not row or row[0] is None:
                continue
            if len(row) < 3:
                raise CommandError("Недостаточно колонок в строке.")
            try:
                position = int(row[0])
            except (ValueError, TypeError):
                raise CommandError(f"Некорректный порядковый номер: {row[0]}")
            category, original_name = row[1], row[2]
            if category not in CATEGORY_SLUGS or not isinstance(original_name, str) or not original_name.strip():
                raise CommandError(f"Ошибка категории или названия в позиции {position}")
            goods.append((position, category, original_name))
        if len(goods) != 373 or {p for p, _, _ in goods} != set(range(1, 374)):
            raise CommandError("Ожидалось ровно 373 уникальные позиции с номерами 1–373.")
        if len({x[2] for x in goods}) != 373:
            raise CommandError("Дословные наименования в файле повторяются; требуется проверка.")
        if options["check_only"]:
            self.stdout.write(self.style.SUCCESS("Проверено: 373 позиции, без записи в базу."))
            return
        with transaction.atomic():
            categories = {}
            for i, (name, slug) in enumerate(CATEGORY_SLUGS.items(), 1):
                category, _ = Category.objects.update_or_create(
                    name=name, defaults={"slug": slug, "order": i},
                )
                categories[name] = category
            for position, name, original in goods:
                TrademarkGood.objects.update_or_create(
                    source_position=position,
                    defaults={"category": categories[name], "original_name": original, "nice_class": 9},
                )
        self.stdout.write(self.style.SUCCESS("Импортировано / обновлено 373 дословные позиции."))
